"""
Comprehensive Futures Analyzer
Analyzes ALL futures assets with buy/sell leg separation and Excel output.
"""

import pandas as pd
import numpy as np
from datetime import datetime, timedelta
import os
import glob
import openpyxl  # For Excel output
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import warnings
warnings.filterwarnings('ignore')

from config import ASSET_CONFIGS, get_config, get_distance_bins
from contract_matching import calculate_buy_sell_distance

# =============================================================================
# CONFIGURATION — all per-asset settings imported from config.py
# =============================================================================

# =============================================================================
# DATA LOADING
# =============================================================================

def load_agent_data(agent_file: str) -> pd.DataFrame:
    """Load agent data and detect buy/sell legs."""
    df = pd.read_csv(agent_file, header=None,
                    names=['date', 'hour', 'min', 'price', 'cum_vol'])

    # Convert Excel date numbers to timestamps
    start_date = datetime(1900, 1, 1)
    df['datetime'] = df.apply(
        lambda row: start_date + timedelta(days=int(row['date']) - 2,
                                          hours=int(row['hour']),
                                          minutes=int(row['min'])),
        axis=1
    )

    df = df.set_index('datetime').sort_index()
    df = df[~df.index.duplicated(keep='last')]

    # Detect position changes (actual trades)
    df['position_change'] = df['cum_vol'].diff()
    df['is_trade'] = df['position_change'].abs() > 0

    # Determine trade direction from position change
    df['leg_type'] = 'flat'
    df.loc[df['position_change'] > 0, 'leg_type'] = 'buy'
    df.loc[df['position_change'] < 0, 'leg_type'] = 'sell'

    return df

def load_contract_csv(path: str) -> pd.DataFrame:
    """Load contract CSV file."""
    df = pd.read_csv(path, header=None)
    df.columns = ["datetime", "open", "high", "low", "close", "volume"]

    df["datetime"] = pd.to_datetime(df["datetime"], format="%Y.%m.%d.%H:%M:%S", errors="coerce")
    df = df.dropna(subset=["datetime"])

    for c in ["open", "high", "low", "close", "volume"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df = df.dropna(subset=["open", "high", "low", "close", "volume"])

    df = df.sort_values("datetime").set_index("datetime")
    df = df[~df.index.duplicated(keep="last")]

    return df

def load_all_contracts(contract_dir: str, exclude_pattern: str = 'aiagent') -> dict:
    """Load all contracts from directory."""
    contracts = {}
    contract_files = sorted(
        f for f in glob.glob(os.path.join(contract_dir, "*.csv"))
        if exclude_pattern.lower() not in os.path.basename(f).lower()
    )

    for file_path in contract_files:
        name = os.path.splitext(os.path.basename(file_path))[0]
        df = load_contract_csv(file_path)
        contracts[name] = df

    return contracts

# =============================================================================
# DISTANCE CALCULATION (with buy/sell separation)
# =============================================================================

# =============================================================================
# ANALYSIS
# =============================================================================

def analyze_asset(asset_name: str, base_dir: str = 'data_export') -> dict:
    """
    Analyze a single asset and return results.
    
    Returns dict with:
    - results_df: DataFrame with all matches
    - stats: Statistics dictionary
    - buy_distribution: Buy leg distance distribution
    - sell_distribution: Sell leg distance distribution
    """
    print(f"\n{'='*70}")
    print(f"Analyzing: {asset_name}")
    print('='*70)

    asset_dir = os.path.join(base_dir, asset_name)

    # Check if directory exists
    if not os.path.exists(asset_dir):
        print(f"  ✗ Directory not found: {asset_dir}")
        return None

    # Get configuration (from centralized config module)
    config = get_config(asset_name)

    # Find agent file
    agent_files = [f for f in os.listdir(asset_dir) if config['agent_pattern'].lower() in f.lower()]
    if not agent_files:
        print(f"  ✗ No agent file found matching pattern: {config['agent_pattern']}")
        return None

    agent_file = os.path.join(asset_dir, agent_files[0])
    print(f"  Agent file: {agent_file}")

    # Load data
    print("  Loading data...")
    agent_df = load_agent_data(agent_file)
    contracts = load_all_contracts(asset_dir)

    print(f"    ✅ Agent data: {len(agent_df):,} bars")
    print(f"    ✅ Contracts: {len(contracts)} loaded")

    # Filter to only trade timestamps (where position actually changed)
    trade_df = agent_df[agent_df['is_trade']].copy()
    print(f"    ✅ Trade timestamps: {len(trade_df):,} (out of {len(agent_df):,} total bars)")

    # Check strict_inside_match mode
    strict_mode = config.get('strict_inside_match', False)
    match_mode_str = "STRICT inside-match (highest volume)" if strict_mode else "min-distance"
    print(f"  Match mode: {match_mode_str}")

    # Analyze matches using 5-minute lookback window
    # For each agent trade at time T, look up contract 1-min bars at T-4, T-3, T-2, T-1, T
    # Composite OHLC: low = min(lows), high = max(highs)
    # Show rollover config
    rollover_dates = config.get('contract_rollover', {})
    if rollover_dates:
        print(f"  Contract rollover dates:")
        for cname, rdate in rollover_dates.items():
            print(f"    {cname}: active from {rdate} onwards")

    print("  Analyzing matches (5-min lookback window)...")
    results = []

    for timestamp, agent_row in trade_df.iterrows():
        agent_price = agent_row['price']
        leg_type = agent_row['leg_type']

        # Build the 5 lookback timestamps: T-4, T-3, T-2, T-1, T
        lookback_times = [timestamp - pd.Timedelta(minutes=i) for i in range(4, -1, -1)]

        # Get rollover config for this asset
        rollover_dates = config.get('contract_rollover', {})

        contract_distances = {}
        for name, cdf in contracts.items():
            # Skip contracts that haven't rolled over yet
            if name in rollover_dates:
                rollover_ts = pd.Timestamp(rollover_dates[name])
                if timestamp < rollover_ts:
                    continue

            # Find which of the 5 lookback timestamps exist in this contract
            existing = cdf.index.intersection(lookback_times)

            if len(existing) == 0:
                contract_distances[name] = None
                continue

            subset = cdf.loc[existing]
            composite_low = subset['low'].min()
            composite_high = subset['high'].max()
            composite_volume = subset['volume'].sum()

            dist_info = calculate_buy_sell_distance(
                agent_price,
                composite_low,
                composite_high,
                leg_type,
                config['tick_size']
            )
            contract_distances[name] = {
                'distance_ticks': dist_info['distance_ticks'],
                'distance_price': dist_info['distance_price'],
                'is_inside': dist_info['is_inside'],
                'reference_point': dist_info['reference_point'],
                'low': composite_low,
                'high': composite_high,
                'volume': composite_volume,
                'bars_in_window': len(existing),
            }

        # Find best match
        valid_matches = {k: v for k, v in contract_distances.items() if v is not None}

        if valid_matches:
            if strict_mode:
                # STRICT MODE: only consider contracts where agent price is INSIDE OHLC range
                inside_matches = {k: v for k, v in valid_matches.items()
                                  if v.get('is_inside', False)}

                if inside_matches:
                    # Among inside matches, pick highest volume (most liquid contract)
                    best_match_name = max(inside_matches, 
                                         key=lambda k: inside_matches[k].get('volume', 0))
                    best_match_info = inside_matches[best_match_name]
                else:
                    # No contract has agent price inside range → unmatched
                    best_match_name = None
                    best_match_info = None
            else:
                # DEFAULT MODE: pick contract with minimum distance (current behavior)
                sorted_matches = sorted(valid_matches.items(),
                                       key=lambda x: x[1]['distance_ticks'])
                best_match_name, best_match_info = sorted_matches[0]

            if best_match_info is not None:
                results.append({
                    'timestamp': timestamp,
                    'agent_price': agent_price,
                    'leg_type': leg_type,
                    'best_match': best_match_name,
                    'best_distance_ticks': best_match_info.get('distance_ticks', 0.0),
                    'best_distance_price': best_match_info.get('distance_price', 0.0),
                    'best_is_inside': best_match_info.get('is_inside', False),
                    'best_low': best_match_info.get('low'),
                    'best_high': best_match_info.get('high'),
                    'best_volume': best_match_info.get('volume'),
                    'bars_in_window': best_match_info.get('bars_in_window'),
                    'contracts_with_data': len(valid_matches)
                })
            else:
                # Strict mode: no inside match found
                results.append({
                    'timestamp': timestamp,
                    'agent_price': agent_price,
                    'leg_type': leg_type,
                    'best_match': None,
                    'best_distance_ticks': None,
                    'best_distance_price': None,
                    'best_is_inside': False,
                    'best_low': None,
                    'best_high': None,
                    'best_volume': None,
                    'bars_in_window': 0,
                    'contracts_with_data': len(valid_matches)  # data existed but no inside match
                })
        else:
            results.append({
                'timestamp': timestamp,
                'agent_price': agent_price,
                'leg_type': leg_type,
                'best_match': None,
                'best_distance_ticks': None,
                'best_distance_price': None,
                'best_is_inside': False,
                'best_low': None,
                'best_high': None,
                'best_volume': None,
                'bars_in_window': 0,
                'contracts_with_data': 0
            })

    df_results = pd.DataFrame(results).set_index('timestamp')

    # Calculate statistics
    stats = calculate_statistics(df_results, asset_name)

    # Calculate buy/sell distributions (using per-asset distance bins)
    distance_bins, bin_labels = get_distance_bins(asset_name)
    buy_dist = calculate_distance_distribution(df_results, 'buy', distance_bins, bin_labels)
    sell_dist = calculate_distance_distribution(df_results, 'sell', distance_bins, bin_labels)

    print(f"  ✅ Analysis complete")

    return {
        'results_df': df_results,
        'stats': stats,
        'buy_distribution': buy_dist,
        'sell_distribution': sell_dist,
        'asset_name': asset_name
    }

def calculate_statistics(df_results: pd.DataFrame, asset_name: str) -> dict:
    """Calculate comprehensive statistics."""
    total = len(df_results)
    with_data = (df_results['contracts_with_data'] > 0).sum()
    inside_range = (df_results['best_is_inside'] == True).sum()
    outside_range = with_data - inside_range

    stats = {
        'asset': asset_name,
        'total_timestamps': total,
        'with_data': with_data,
        'inside_ohlc': inside_range,
        'outside_ohlc': outside_range,
        'coverage_pct': (with_data / total * 100) if total > 0 else 0,
        'inside_pct': (inside_range / total * 100) if total > 0 else 0,
        'outside_pct': (outside_range / total * 100) if total > 0 else 0
    }

    # Buy leg statistics
    buy_df = df_results[df_results['leg_type'] == 'buy']
    if len(buy_df) > 0:
        buy_count = len(buy_df)
        buy_inside = (buy_df['best_is_inside'] == True).sum()
        buy_outside = buy_count - buy_inside
        buy_outside_df = buy_df[~buy_df['best_is_inside']]

        stats['buy_count'] = buy_count
        stats['buy_inside'] = buy_inside
        stats['buy_outside'] = buy_outside
        stats['buy_inside_pct'] = (buy_inside / buy_count * 100) if buy_count > 0 else 0
        stats['buy_outside_pct'] = (buy_outside / buy_count * 100) if buy_count > 0 else 0

        if buy_outside > 0:
            stats['buy_avg_distance'] = buy_outside_df['best_distance_ticks'].mean()
            stats['buy_median_distance'] = buy_outside_df['best_distance_ticks'].median()
            stats['buy_max_distance'] = buy_outside_df['best_distance_ticks'].max()

    # Sell leg statistics
    sell_df = df_results[df_results['leg_type'] == 'sell']
    if len(sell_df) > 0:
        sell_count = len(sell_df)
        sell_inside = (sell_df['best_is_inside'] == True).sum()
        sell_outside = sell_count - sell_inside
        sell_outside_df = sell_df[~sell_df['best_is_inside']]

        stats['sell_count'] = sell_count
        stats['sell_inside'] = sell_inside
        stats['sell_outside'] = sell_outside
        stats['sell_inside_pct'] = (sell_inside / sell_count * 100) if sell_count > 0 else 0
        stats['sell_outside_pct'] = (sell_outside / sell_count * 100) if sell_count > 0 else 0

        if sell_outside > 0:
            stats['sell_avg_distance'] = sell_outside_df['best_distance_ticks'].mean()
            stats['sell_median_distance'] = sell_outside_df['best_distance_ticks'].median()
            stats['sell_max_distance'] = sell_outside_df['best_distance_ticks'].max()

    return stats

def calculate_distance_distribution(df_results: pd.DataFrame, leg_type: str,
                                    distance_bins: list, bin_labels: list) -> pd.DataFrame:
    """Calculate distance distribution for a specific leg type."""
    leg_df = df_results[df_results['leg_type'] == leg_type]

    if len(leg_df) == 0:
        return pd.DataFrame()

    distances = leg_df['best_distance_ticks'].fillna(0)
    counts, _ = np.histogram(distances, bins=distance_bins)

    df_dist = pd.DataFrame({
        'distance_range': bin_labels,
        'count': counts,
        'percentage': (counts / len(leg_df) * 100) if len(leg_df) > 0 else 0
    })

    return df_dist

# =============================================================================
# EXCEL OUTPUT
# =============================================================================

def create_excel_workbook(all_results: dict, output_file: str = 'all_futures_analysis.xlsx'):
    """Create Excel workbook with all results."""
    print(f"\n{'='*70}")
    print("Creating Excel workbook...")
    print('='*70)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # 1. Summary Sheet
    create_summary_sheet(wb, all_results)

    # 2. Individual asset sheets
    for asset_name, data in all_results.items():
        if data is None:
            continue

        # Create main sheet
        sheet_name = asset_name[:31]  # Excel sheet name limit
        create_asset_sheet(wb, sheet_name, data)

        # Create buy distribution sheet
        buy_sheet_name = f"{sheet_name[:27]}_Buy"
        create_distribution_sheet(wb, buy_sheet_name, data['buy_distribution'], 'Buy')

        # Create sell distribution sheet
        sell_sheet_name = f"{sheet_name[:26]}_Sell"
        create_distribution_sheet(wb, sell_sheet_name, data['sell_distribution'], 'Sell')

    # Save workbook
    wb.save(output_file)
    print(f"  ✅ Saved: {output_file}")
    print(f"  Sheets: {len(wb.sheetnames)}")

    return output_file

def create_summary_sheet(wb, all_results: dict):
    """Create summary sheet with all asset statistics."""
    ws = wb.create_sheet('Summary', 0)

    # Collect all stats
    summary_data = []
    for asset_name, data in all_results.items():
        if data is None:
            continue
        stats = data['stats']
        summary_data.append(stats)

    df_summary = pd.DataFrame(summary_data)

    # Write to sheet
    for r_idx, row in enumerate(dataframe_to_rows(df_summary, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Format header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Auto-fit columns
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width

    print(f"  ✅ Created: Summary sheet")

def create_asset_sheet(wb, sheet_name: str, data: dict):
    """Create sheet with detailed results for an asset."""
    ws = wb.create_sheet(sheet_name)

    # Write results dataframe
    df = data['results_df'].reset_index()
    df['timestamp'] = df['timestamp'].dt.strftime('%Y-%m-%d %H:%M:%S')

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Format header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Auto-fit columns
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[column].width = adjusted_width

    print(f"  ✅ Created: {sheet_name} sheet ({len(df)} rows)")

def create_distribution_sheet(wb, sheet_name: str, df_dist: pd.DataFrame, leg_type: str):
    """Create sheet with distance distribution for a leg type."""
    ws = wb.create_sheet(sheet_name)

    if len(df_dist) == 0:
        ws.cell(row=1, column=1, value=f"No {leg_type} legs found")
        return

    # Write dataframe
    for r_idx, row in enumerate(dataframe_to_rows(df_dist, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Format header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Format percentage column
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right')

    # Auto-fit columns
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 20)
        ws.column_dimensions[column].width = adjusted_width

    print(f"  ✅ Created: {sheet_name} sheet ({leg_type} distribution)")

# =============================================================================
# MAIN PIPELINE
# =============================================================================

def analyze_all_futures():
    """Analyze all futures assets and create Excel report."""
    print("="*70)
    print("COMPREHENSIVE FUTURES ANALYSIS")
    print("Analyzing all assets with buy/sell leg separation")
    print("="*70)

    # Get all asset directories
    base_dir = 'data_export'
    asset_dirs = [d for d in os.listdir(base_dir)
                  if os.path.isdir(os.path.join(base_dir, d)) and d != '.']

    print(f"\nFound {len(asset_dirs)} asset directories:")
    for asset in sorted(asset_dirs):
        print(f"  - {asset}")

    # Analyze each asset
    all_results = {}
    for asset_name in sorted(asset_dirs):
        result = analyze_asset(asset_name, base_dir)
        all_results[asset_name] = result

    # Create Excel workbook
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_file = f'all_futures_analysis_{timestamp}.xlsx'
    create_excel_workbook(all_results, output_file)

    # Print final summary
    print("\n" + "="*70)
    print("ANALYSIS COMPLETE")
    print("="*70)

    print(f"\nResults saved to: {output_file}")
    print(f"\nSummary Statistics:")
    print("-" * 70)

    for asset_name, data in sorted(all_results.items()):
        if data is None:
            continue

        stats = data['stats']
        print(f"\n{asset_name}:")
        print(f"  Total:       {stats['total_timestamps']:,}")
        print(f"  Coverage:    {stats['with_data']:,} ({stats['coverage_pct']:.1f}%)")
        print(f"  Inside OHLC: {stats['inside_ohlc']:,} ({stats['inside_pct']:.1f}%)")
        print(f"  Outside OHLC: {stats['outside_ohlc']:,} ({stats['outside_pct']:.1f}%)")

        if 'buy_count' in stats:
            print(f"  Buy legs:    {stats['buy_count']:,} (inside: {stats['buy_inside']:,} {stats['buy_inside_pct']:.1f}%)")
        if 'sell_count' in stats:
            print(f"  Sell legs:   {stats['sell_count']:,} (inside: {stats['sell_inside']:,} {stats['sell_inside_pct']:.1f}%)")

    print("\n" + "="*70)

    return all_results

if __name__ == '__main__':
    all_results = analyze_all_futures()
